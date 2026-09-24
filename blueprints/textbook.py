from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file, make_response, jsonify, session
from database import db, Class, Institution, SchoolYear, TextbookSession, CourseModule, CourseSection, Objective, Competency, AppSetting, TimeSlot, ScheduleEntry, ExceptionalEvent
from datetime import datetime
import io
import socket
import secrets
import qrcode
import qrcode.constants

def get_local_ip():
    """Détecte l'adresse IP locale de la machine sur le réseau local."""
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.settimeout(0.2)
        s.connect(('10.255.255.255', 1))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        try:
            return socket.gethostbyname(socket.gethostname())
        except Exception:
            return '127.0.0.1'

def get_current_school_year():
    """Récupère l'année scolaire active pour la session courante."""
    year_id = session.get('active_school_year_id')
    if year_id:
        y = SchoolYear.query.get(year_id)
        if y: return y
    active = SchoolYear.query.filter_by(is_active=True).first()
    if active: return active
    return SchoolYear.query.order_by(SchoolYear.start_date.desc()).first()

# WeasyPrint for PDF generation
HTML = None  # Will be set if weasyprint is available
try:
    from weasyprint import HTML  # type: ignore[import-untyped]
    weasyprint_available = True
except ImportError:
    weasyprint_available = False
except OSError:
    # Handle missing shared libraries (like libpango/libcairo) gently
    weasyprint_available = False

textbook_bp = Blueprint('textbook', __name__)

def seed_textbook_data():
    """Helper to ensure at least one Institution and SchoolYear exist."""
    if not Institution.query.first():
        inst = Institution(name="Lycée d'Excellence", address="123 Rue de l'École", city="Casablanca")
        db.session.add(inst)
        db.session.commit()
        
        # Default Year
        year = SchoolYear(
            name="2025-2026", 
            start_date=datetime(2025, 9, 1).date(), 
            end_date=datetime(2026, 6, 30).date(),
            institution_id=inst.id
        )
        db.session.add(year)
        db.session.commit()

@textbook_bp.route('/textbook')
def index():
    seed_textbook_data() # Ensure data exists
    
    institution = Institution.query.first()
    years = SchoolYear.query.filter_by(institution_id=institution.id).all() if institution else []
    
    current_year_active = get_current_school_year()
    if current_year_active:
        classes = Class.query.filter_by(school_year_id=current_year_active.id).order_by(Class.name).all()
    else:
        classes = Class.query.order_by(Class.name).all()
    
    # Pre-fetch M/S for filters if needed, but for list view we just need to render sessions
    # Sessions now have relations so session.module.name works
    
    # Filters
    year_id = request.args.get('year_id')
    class_id = request.args.get('class_id')
    search_q = request.args.get('q', '').strip()
    
    query = TextbookSession.query
    
    current_year = None
    if year_id:
        query = query.filter_by(school_year_id=year_id)
        current_year = db.session.get(SchoolYear, year_id)
    else:
        current_year = get_current_school_year()
        if current_year:
            query = query.filter_by(school_year_id=current_year.id)
            
    if class_id:
        query = query.filter_by(class_id=class_id)

    if search_q:
        search_term = f"%{search_q}%"
        query = query.join(CourseModule).join(CourseSection).filter(
            (TextbookSession.course_titles.ilike(search_term)) |
            (CourseModule.name.ilike(search_term)) |
            (CourseSection.name.ilike(search_term)) |
            (TextbookSession.teacher_name.ilike(search_term)) |
            (TextbookSession.remark.ilike(search_term))
        )
        
    sessions = query.order_by(TextbookSession.date.desc(), TextbookSession.id.desc()).all()

    # KPI Statistics for the dashboard
    year_sessions_query = TextbookSession.query
    if current_year:
        year_sessions_query = year_sessions_query.filter_by(school_year_id=current_year.id)
    all_year_sessions = year_sessions_query.all()

    total_sessions_count = len(all_year_sessions)
    now = datetime.now()
    current_month_sessions_count = sum(1 for s in all_year_sessions if s.date.year == now.year and s.date.month == now.month)
    distinct_classes_count = len(set(s.class_id for s in all_year_sessions))
    distinct_modules_count = len(set(s.module_id for s in all_year_sessions))
    total_hours_count = round(sum((s.duration_hours or 2.0) for s in all_year_sessions), 1)

    stats = {
        'total_sessions': total_sessions_count,
        'month_sessions': current_month_sessions_count,
        'active_classes': distinct_classes_count,
        'covered_modules': distinct_modules_count,
        'total_hours': total_hours_count
    }

    # Calcul de progression pédagogique par module (% d'objectifs couverts)
    all_modules = CourseModule.query.all()
    # Séances du filtre actif (année + classe éventuelle)
    scope_sessions = sessions if class_id else all_year_sessions
    covered_obj_ids = set()
    for s in scope_sessions:
        for obj in s.objectives:
            covered_obj_ids.add(obj.id)

    module_progress_list = []
    for m in all_modules:
        # Tous les objectifs du module via ses chapitres (sections) ou ses compétences
        module_objs_dict = {}
        for sec in m.sections:
            for obj in sec.objectives:
                module_objs_dict[obj.id] = obj
        for comp in m.competencies:
            for obj in comp.objectives:
                module_objs_dict[obj.id] = obj
        
        total_objs_in_module = list(module_objs_dict.values())
        total_count = len(total_objs_in_module)
        covered_count = sum(1 for o in total_objs_in_module if o.id in covered_obj_ids)
        percent = round((covered_count / total_count * 100)) if total_count > 0 else 0
        
        # Heures consacrées au module
        module_hours = round(sum((s.duration_hours or 2.0) for s in scope_sessions if s.module_id == m.id), 1)

        module_progress_list.append({
            'module': m,
            'total_objectives': total_count,
            'covered_objectives': covered_count,
            'percent': percent,
            'hours': module_hours
        })
    
    # Données de partage sécurisé (Option 3 - Jeton Magic Link)
    share_enabled = AppSetting.get_value('textbook_share_enabled', 'false').lower() == 'true'
    share_token = AppSetting.get_value('textbook_share_token', '')
    if not share_token:
        share_token = secrets.token_urlsafe(16)
        AppSetting.set_value('textbook_share_token', share_token, "Jeton secret d'accès au cahier de texte partagé")

    local_ip = get_local_ip()
    port = request.host.split(':')[1] if ':' in request.host else '5000'
    base_url = f"http://{local_ip}:{port}"
    share_url = f"{base_url}/partage/cahier-texte?token={share_token}"
    if class_id:
        share_url += f"&class_id={class_id}"

    # Récupération des cours prévus aujourd'hui selon l'emploi du temps
    today = datetime.now().date()
    today_weekday = today.weekday() # 0 = Lundi ... 5 = Samedi
    today_schedule = []
    if current_year and today_weekday < 6:
        entries = ScheduleEntry.query.filter_by(
            school_year_id=current_year.id,
            day_of_week=today_weekday
        ).all()
        for e in entries:
            # Vérifier si une séance a déjà été enregistrée pour cette classe à cette date
            has_session = TextbookSession.query.filter_by(
                class_id=e.class_id,
                date=today,
                school_year_id=current_year.id
            ).first() is not None
            
            today_schedule.append({
                'entry': e,
                'has_session': has_session,
                'class_id': e.class_id,
                'class_name': e.classe.name if e.classe else '',
                'slot': e.time_slot,
                'room': e.room,
                'subject': e.subject_title
            })

    # Événements exceptionnels de l'année (Grèves, Événements établissement, Séances suspendues)
    exceptional_events_query = ExceptionalEvent.query
    if current_year:
        exceptional_events_query = exceptional_events_query.filter_by(school_year_id=current_year.id)
    if class_id:
        exceptional_events_query = exceptional_events_query.filter(
            (ExceptionalEvent.class_id == class_id) | (ExceptionalEvent.class_id.is_(None))
        )
    exceptional_events = exceptional_events_query.order_by(ExceptionalEvent.date.desc()).all()

    return render_template('textbook_list.html', 
                           institution=institution, 
                           years=years, 
                           classes=classes, 
                           sessions=sessions,
                           current_year=current_year,
                           today_schedule=today_schedule,
                           today_date_str=today.strftime('%d/%m/%Y'),
                           stats=stats,
                           module_progress=module_progress_list,
                           search_q=search_q,
                           selected_class_id=int(class_id) if class_id else None,
                           share_enabled=share_enabled,
                           share_token=share_token,
                           share_url=share_url,
                           exceptional_events=exceptional_events)

@textbook_bp.route('/textbook/new', methods=['GET', 'POST'])
def new_session():
    current_year = get_current_school_year()
    if current_year and current_year.is_read_only():
        flash(f"L'année scolaire {current_year.name} est archivée en lecture seule. Impossible d'ajouter de nouvelles séances.", 'error')
        return redirect(url_for('textbook.index'))

    duplicate_id = request.args.get('duplicate_from')
    source_session = None
    if duplicate_id:
        source_session = TextbookSession.query.get(duplicate_id)

    if request.method == 'POST':
        try:
            module_id = request.form.get('module_id')
            section_id = request.form.get('section_id')
            school_year_id = request.form.get('school_year_id')
            # Support de sélection multi-classes : 'class_ids' (liste de checkboxes) ou fallback 'class_id'
            class_ids = request.form.getlist('class_ids')
            if not class_ids:
                single_class = request.form.get('class_id')
                if single_class:
                    class_ids = [single_class]

            if not all([module_id, section_id, class_ids, school_year_id]):
                flash('Veuillez remplir tous les champs obligatoires (Module, Section, au moins une Classe, Année).', 'error')
                return redirect(url_for('textbook.new_session'))

            obj_ids = request.form.getlist('objective_ids')
            comp_ids = request.form.getlist('competency_ids')
            session_date = datetime.strptime(request.form.get('date') or '', '%Y-%m-%d').date()
            titles = request.form.get('course_titles') or ''
            remark_val = request.form.get('remark')
            t_name = request.form.get('teacher_name') or session.get('user_full_name') or session.get('username') or ''
            s_type = request.form.get('session_type') or 'Cours'
            hw_val = request.form.get('homework') or None
            hw_due_str = request.form.get('homework_due_date')
            hw_due = datetime.strptime(hw_due_str, '%Y-%m-%d').date() if hw_due_str else None
            s_time = request.form.get('start_time') or '08:30'
            e_time = request.form.get('end_time') or '10:30'
            
            # Calcul automatique du volume horaire
            try:
                t1 = datetime.strptime(s_time, '%H:%M')
                t2 = datetime.strptime(e_time, '%H:%M')
                diff_hours = (t2 - t1).total_seconds() / 3600.0
                dur_hours = round(diff_hours, 2) if diff_hours > 0 else 2.0
            except Exception:
                dur_hours = 2.0

            created_count = 0
            for cid_str in class_ids:
                cid = int(cid_str)
                new_sess = TextbookSession(
                    module_id=int(module_id),  # type: ignore[arg-type]
                    section_id=int(section_id),  # type: ignore[arg-type]
                    course_titles=titles,
                    remark=remark_val,
                    date=session_date,
                    class_id=cid,
                    school_year_id=int(school_year_id),  # type: ignore[arg-type]
                    teacher_name=t_name,
                    session_type=s_type,
                    homework=hw_val,
                    homework_due_date=hw_due,
                    start_time=s_time,
                    end_time=e_time,
                    duration_hours=dur_hours
                )
                db.session.add(new_sess)

                # Attacher objectifs & compétences
                for oid in obj_ids:
                    obj = db.session.get(Objective, int(oid))
                    if obj: new_sess.objectives.append(obj)

                for comp_id in comp_ids:
                    comp = db.session.get(Competency, int(comp_id))
                    if comp: new_sess.competencies.append(comp)

                created_count += 1

            db.session.commit()
            if created_count > 1:
                flash(f'{created_count} séances créées avec succès pour les classes sélectionnées.', 'success')
            else:
                flash('Séance ajoutée avec succès.', 'success')
            return redirect(url_for('textbook.index'))
        except Exception as e:
            flash(f'Erreur: {str(e)}', 'error')
    
    # Data for form
    institution = Institution.query.first()
    years = SchoolYear.query.filter_by(institution_id=institution.id).all() if institution else []
    current_year_active = get_current_school_year()
    if current_year_active:
        classes = Class.query.filter_by(school_year_id=current_year_active.id).order_by(Class.name).all()
    else:
        classes = Class.query.order_by(Class.name).all()
    modules = CourseModule.query.all()
    time_slots = TimeSlot.query.filter_by(is_active=True).order_by(TimeSlot.order_num, TimeSlot.start_time).all()
    default_teacher = session.get('user_full_name') or session.get('username') or ''
    default_date = datetime.now().strftime('%Y-%m-%d')
    
    # Préremplissage depuis le planning ou la feuille d'appel (URL query params)
    prefill_class_id = request.args.get('class_id', type=int)
    prefill_start_time = request.args.get('start_time')
    prefill_end_time = request.args.get('end_time')
    prefill_date_str = request.args.get('date')
    prefill_session_type = request.args.get('session_type')
    prefill_remark = request.args.get('remark')
    prefill_title = request.args.get('title')

    if prefill_date_str:
        default_date = prefill_date_str

    if not source_session and (prefill_class_id or prefill_start_time or prefill_end_time or prefill_date_str or prefill_remark):
        class DummySession:
            id = None
        ds = DummySession()
        ds.id = None
        ds.class_id = prefill_class_id
        ds.school_year_id = current_year_active.id if current_year_active else None
        ds.start_time = prefill_start_time or '08:30'
        ds.end_time = prefill_end_time or '10:30'
        ds.duration_hours = 2.0
        try:
            ds.date = datetime.strptime(prefill_date_str, '%Y-%m-%d').date() if prefill_date_str else datetime.now().date()
        except Exception:
            ds.date = datetime.now().date()
        ds.session_type = prefill_session_type or 'Cours'
        ds.course_titles = prefill_title or ''
        ds.remark = prefill_remark or ''
        ds.homework = ''
        ds.homework_due_date = None
        ds.module_id = None
        ds.section_id = None
        ds.objectives = []
        ds.competencies = []
        ds.teacher_name = default_teacher
        source_session = ds


    return render_template('textbook_form.html', 
                           years=years, 
                           classes=classes, 
                           modules=modules,
                           time_slots=time_slots,
                           textbook_session=source_session,
                           default_teacher=default_teacher,
                           default_date=default_date,
                           is_duplicating=bool(duplicate_id))

@textbook_bp.route('/textbook/<int:id>/edit', methods=['GET', 'POST'])
def edit_session(id):
    session_record = TextbookSession.query.get_or_404(id)
    
    if request.method == 'POST':
        current_year = get_current_school_year()
        if (current_year and current_year.is_read_only()) or (session_record.year and session_record.year.is_read_only()):
            flash("Cette séance appartient à une année scolaire archivée en lecture seule.", 'error')
            return redirect(url_for('textbook.index'))

        try:
            module_id = request.form.get('module_id')
            section_id = request.form.get('section_id')
            class_id = request.form.get('class_id')
            school_year_id = request.form.get('school_year_id')

            if not all([module_id, section_id, class_id, school_year_id]):
                flash('Champs obligatoires manquants.', 'error')
                return redirect(url_for('textbook.edit_session', id=id))

            # After the `not all(...)` guard above, these are guaranteed to be str
            session_record.module_id = int(module_id)  # type: ignore[arg-type]
            session_record.section_id = int(section_id)  # type: ignore[arg-type]
            session_record.course_titles = request.form.get('course_titles') or ''
            session_record.remark = request.form.get('remark')
            session_record.date = datetime.strptime(request.form.get('date') or '', '%Y-%m-%d').date()
            session_record.class_id = int(class_id)  # type: ignore[arg-type]
            session_record.school_year_id = int(school_year_id)  # type: ignore[arg-type]
            session_record.teacher_name = request.form.get('teacher_name')
            session_record.session_type = request.form.get('session_type') or 'Cours'
            session_record.homework = request.form.get('homework') or None
            hw_due_str = request.form.get('homework_due_date')
            session_record.homework_due_date = datetime.strptime(hw_due_str, '%Y-%m-%d').date() if hw_due_str else None
            session_record.start_time = request.form.get('start_time') or '08:30'
            session_record.end_time = request.form.get('end_time') or '10:30'

            try:
                t1 = datetime.strptime(session_record.start_time, '%H:%M')
                t2 = datetime.strptime(session_record.end_time, '%H:%M')
                diff_hours = (t2 - t1).total_seconds() / 3600.0
                session_record.duration_hours = round(diff_hours, 2) if diff_hours > 0 else 2.0
            except Exception:
                session_record.duration_hours = 2.0
            
            # Update M2M
            session_record.objectives = []
            session_record.competencies = []
            
            obj_ids = request.form.getlist('objective_ids')
            comp_ids = request.form.getlist('competency_ids')
            
            for oid in obj_ids:
                obj = db.session.get(Objective, int(oid))
                if obj: session_record.objectives.append(obj)
                
            for cid in comp_ids:
                comp = db.session.get(Competency, int(cid))
                if comp: session_record.competencies.append(comp)
            
            db.session.commit()
            flash('Séance modifiée avec succès.', 'success')
            return redirect(url_for('textbook.index'))
        except Exception as e:
            flash(f'Erreur: {str(e)}', 'error')

    institution = Institution.query.first()
    years = SchoolYear.query.filter_by(institution_id=institution.id).all() if institution else []
    cur_year = get_current_school_year()
    target_year_id = session_record.school_year_id or (cur_year.id if cur_year else None)
    if target_year_id:
        classes = Class.query.filter_by(school_year_id=target_year_id).order_by(Class.name).all()
    else:
        classes = Class.query.order_by(Class.name).all()
    modules = CourseModule.query.all()
    time_slots = TimeSlot.query.filter_by(is_active=True).order_by(TimeSlot.order_num, TimeSlot.start_time).all()
    
    return render_template('textbook_form.html', 
                           years=years, 
                           classes=classes, 
                           modules=modules, 
                           time_slots=time_slots,
                           textbook_session=session_record)

@textbook_bp.route('/api/module/<int:module_id>/sections')
def get_module_sections(module_id):
    sections = CourseSection.query.filter_by(module_id=module_id).all()
    return jsonify([{'id': s.id, 'name': s.name} for s in sections])

@textbook_bp.route('/api/section/<int:section_id>/objectives')
def get_section_objectives(section_id):
    objectives = Objective.query.filter_by(section_id=section_id).all()
    return jsonify([{'id': o.id, 'description': o.description, 'section_id': o.section_id} for o in objectives])

@textbook_bp.route('/api/module/<int:module_id>/competencies')
def get_module_competencies(module_id):
    competencies = Competency.query.filter_by(module_id=module_id).all()
    return jsonify([{'id': c.id, 'description': c.description} for c in competencies])

@textbook_bp.route('/api/competencies/objectives', methods=['POST'])
def get_competencies_objectives():
    comp_ids = request.json.get('competency_ids', [])
    if not comp_ids:
        return jsonify([])
    
    objectives = Objective.query.filter(Objective.competency_id.in_(comp_ids)).all()
    return jsonify([{'id': o.id, 'description': o.description, 'competency_id': o.competency_id} for o in objectives])

@textbook_bp.route('/api/textbook/events')
def api_events():
    """Retourne les séances au format FullCalendar JSON."""
    year_id = request.args.get('year_id')
    class_id = request.args.get('class_id')
    
    query = TextbookSession.query
    if year_id:
        query = query.filter_by(school_year_id=year_id)
    else:
        current_year = get_current_school_year()
        if current_year:
            query = query.filter_by(school_year_id=current_year.id)
            
    if class_id:
        query = query.filter_by(class_id=class_id)
        
    sessions = query.all()
    
    type_colors = {
        'Cours': {'bg': '#4e73df', 'border': '#2e59d9'},
        'TD': {'bg': '#f6c23e', 'border': '#dda20a'},
        'TP': {'bg': '#1cc88a', 'border': '#17a673'},
        'Évaluation': {'bg': '#e74a3b', 'border': '#be2617'},
        'Projet': {'bg': '#36b9cc', 'border': '#258391'},
        'Soutien': {'bg': '#6c757d', 'border': '#545b62'}
    }
    
    events = []
    for s in sessions:
        date_str = s.date.strftime('%Y-%m-%d')
        start_str = f"{date_str}T{s.start_time or '08:30'}:00"
        end_str = f"{date_str}T{s.end_time or '10:30'}:00"
        
        stype = s.session_type or 'Cours'
        color_info = type_colors.get(stype, {'bg': '#4e73df', 'border': '#2e59d9'})
        
        title = f"[{s.classe.name}] {stype}: {s.course_titles}" if s.classe else f"{stype}: {s.course_titles}"
        if len(title) > 40:
            title = title[:37] + "..."
            
        events.append({
            'id': s.id,
            'title': title,
            'start': start_str,
            'end': end_str,
            'backgroundColor': color_info['bg'],
            'borderColor': color_info['border'],
            'textColor': '#000000' if stype == 'TD' else '#ffffff',
            'extendedProps': {
                'className': s.classe.name if s.classe else '',
                'classId': s.class_id,
                'moduleName': s.module.name if s.module else '',
                'sectionName': s.section.name if s.section else '',
                'sessionType': stype,
                'courseTitles': s.course_titles,
                'homework': s.homework or '',
                'homeworkDueDate': s.homework_due_date.strftime('%d/%m/%Y') if s.homework_due_date else '',
                'remark': s.remark or '',
                'teacher': s.teacher_name or '',
                'dateFormatted': s.date.strftime('%d/%m/%Y'),
                'timeSlot': f"{s.start_time or '08:30'} - {s.end_time or '10:30'}",
                'editUrl': url_for('textbook.edit_session', id=s.id),
                'pdfUrl': url_for('textbook.export_pdf', id=s.id),
                'attendanceUrl': url_for('absence.index', class_id=s.class_id, date=s.date.strftime('%Y-%m-%d'))
            }
        })
        
    return jsonify(events)

# =========================================================================
# ROUTES DU LIEN MAGIQUE DE PARTAGE (OPTION 3)
# =========================================================================

@textbook_bp.route('/textbook/share/toggle', methods=['POST'])
def toggle_share():
    """Active ou désactive le partage externe du cahier de texte."""
    current_status = AppSetting.get_value('textbook_share_enabled', 'false').lower() == 'true'
    new_status = not current_status
    AppSetting.set_value('textbook_share_enabled', 'true' if new_status else 'false', "Statut du partage du cahier de texte")
    
    # S'assurer qu'un token existe
    token = AppSetting.get_value('textbook_share_token', '')
    if not token:
        token = secrets.token_urlsafe(16)
        AppSetting.set_value('textbook_share_token', token, "Jeton secret d'accès au cahier de texte")

    local_ip = get_local_ip()
    port = request.host.split(':')[1] if ':' in request.host else '5000'
    share_url = f"http://{local_ip}:{port}/partage/cahier-texte?token={token}"

    return jsonify({
        'success': True,
        'enabled': new_status,
        'token': token,
        'share_url': share_url,
        'message': "Partage ACTIVÉ avec succès" if new_status else "Partage DÉSACTIVÉ"
    })

@textbook_bp.route('/textbook/share/revoke-token', methods=['POST'])
def revoke_share_token():
    """Régénère un nouveau jeton d'accès secret (invalide tous les anciens liens)."""
    new_token = secrets.token_urlsafe(16)
    AppSetting.set_value('textbook_share_token', new_token, "Jeton secret d'accès au cahier de texte")
    
    local_ip = get_local_ip()
    port = request.host.split(':')[1] if ':' in request.host else '5000'
    share_url = f"http://{local_ip}:{port}/partage/cahier-texte?token={new_token}"

    return jsonify({
        'success': True,
        'token': new_token,
        'share_url': share_url,
        'message': "Nouveau lien de partage généré. L'ancien lien est révoqué."
    })

@textbook_bp.route('/textbook/share/qr')
def shared_qr():
    """Génère le QR Code du lien de partage pour scan direct."""
    token = AppSetting.get_value('textbook_share_token', '')
    class_id = request.args.get('class_id')
    local_ip = get_local_ip()
    port = request.host.split(':')[1] if ':' in request.host else '5000'
    
    url = f"http://{local_ip}:{port}/partage/cahier-texte?token={token}"
    if class_id:
        url += f"&class_id={class_id}"

    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=3,
    )
    qr.add_data(url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="#1e293b", back_color="#ffffff")

    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return send_file(buf, mimetype='image/png')

@textbook_bp.route('/partage/cahier-texte')
def shared_view():
    """Vue publique officielle accessible par lien magique avec jeton secret."""
    is_enabled = AppSetting.get_value('textbook_share_enabled', 'false').lower() == 'true'
    token_provided = request.args.get('token', '').strip()
    valid_token = AppSetting.get_value('textbook_share_token', '')

    # Vérification du statut ON/OFF et du jeton
    if not is_enabled:
        return render_template('textbook_shared_disabled.html', 
                               reason="Le partage du cahier de texte est actuellement désactivé par l'enseignant.")
        
    if not token_provided or token_provided != valid_token:
        return render_template('textbook_shared_disabled.html', 
                               reason="Lien invalide ou jeton expiré / révoqué.")

    institution = Institution.query.first() or Institution(name="Établissement Scolaire")
    current_year = get_current_school_year()
    classes = Class.query.filter_by(school_year_id=current_year.id).order_by(Class.name).all() if current_year else Class.query.order_by(Class.name).all()

    class_id = request.args.get('class_id')
    selected_class = None
    query = TextbookSession.query
    if current_year:
        query = query.filter_by(school_year_id=current_year.id)
    if class_id:
        query = query.filter_by(class_id=class_id)
        selected_class = db.session.get(Class, class_id)

    sessions = query.order_by(TextbookSession.date.desc(), TextbookSession.id.desc()).all()

    # Statistiques du partage
    total_sessions = len(sessions)
    total_hours = round(sum((s.duration_hours or 2.0) for s in sessions), 1)

    # Calcul de la progression pédagogique
    modules = CourseModule.query.all()
    covered_obj_ids = set()
    for s in sessions:
        for obj in s.objectives:
            covered_obj_ids.add(obj.id)

    module_progress_list = []
    for m in modules:
        total_objs_in_module = []
        for comp in m.competencies:
            total_objs_in_module.extend(comp.objectives)
        total_count = len(total_objs_in_module)
        covered_count = sum(1 for o in total_objs_in_module if o.id in covered_obj_ids)
        pct = round((covered_count / total_count * 100)) if total_count > 0 else 0
        module_hours = round(sum((s.duration_hours or 2.0) for s in sessions if s.module_id == m.id), 1)
        module_progress_list.append({
            'module': m,
            'total_objectives': total_count,
            'covered_objectives': covered_count,
            'percent': pct,
            'hours': module_hours
        })

    # Nom de l'enseignant déduit
    teachers = list(filter(None, [s.teacher_name for s in sessions]))
    teacher_name = teachers[0] if teachers else "Enseignant"

    return render_template('textbook_shared_view.html',
                           institution=institution,
                           school_year=current_year,
                           classes=classes,
                           selected_class=selected_class,
                           selected_class_id=int(class_id) if class_id else None,
                           sessions=sessions,
                           total_sessions=total_sessions,
                           total_hours=total_hours,
                           module_progress=module_progress_list,
                           teacher_name=teacher_name,
                           token=token_provided,
                           now=datetime.now())

@textbook_bp.route('/textbook/<int:id>/pdf')
def export_pdf(id):
    session = TextbookSession.query.get_or_404(id)
    html = render_template('textbook_pdf_template.html', session=session, title="Fiche de Séance")
    return generate_pdf(html, f"seance_{session.id}.pdf")

@textbook_bp.route('/textbook/progression')
def progression_report():
    year_id = request.args.get('year_id')
    institution = Institution.query.first()
    years = SchoolYear.query.filter_by(institution_id=institution.id).all() if institution else []
    
    current_year = None
    if year_id:
        current_year = db.session.get(SchoolYear, int(year_id))
    elif years:
        current_year = years[-1]

    if not current_year:
        flash("Aucune année scolaire trouvée.", "warning")
        return redirect(url_for('textbook.index'))

    sessions = TextbookSession.query.filter_by(school_year_id=current_year.id).order_by(TextbookSession.date).all()
    
    # Grouping logic
    # Key: (module_id, section_id, course_titles, objectives_id_tuple)
    grouped_data = {}
    
    for s in sessions:
        obj_ids = tuple(sorted([o.id for o in s.objectives]))
        # Normalize strings for better grouping
        titles_norm = s.course_titles.strip() if s.course_titles else ""
        key = (s.module_id, s.section_id, titles_norm, obj_ids)
        
        if key not in grouped_data:
            grouped_data[key] = {
                'module': s.module.name,
                'section': s.section.name,
                'titles': s.course_titles,
                'objectives': [o.description for o in s.objectives],
                'entries': [] # List of {class: name, date: date}
            }
        
        grouped_data[key]['entries'].append({
            'class': s.classe.name,
            'date': s.date
        })

    # Convert to list and sort entries by class name within each group if needed
    report_items = list(grouped_data.values())

    return render_template('textbook_progress.html', 
                           report_items=report_items, 
                           current_year=current_year,
                           years=years,
                           institution=institution)

@textbook_bp.route('/textbook/search')
def search_sessions():
    # Filters
    class_id = request.args.get('class_id')
    module_id = request.args.get('module_id')
    section_id = request.args.get('section_id')
    
    query = TextbookSession.query
    
    if class_id:
        query = query.filter_by(class_id=class_id)
    if module_id:
        query = query.filter_by(module_id=module_id)
    if section_id:
        query = query.filter_by(section_id=section_id)
        
    sessions = query.order_by(TextbookSession.date.desc()).all()
    
    # Data for filters
    current_year_active = get_current_school_year()
    if current_year_active:
        classes = Class.query.filter_by(school_year_id=current_year_active.id).order_by(Class.name).all()
    else:
        classes = Class.query.order_by(Class.name).all()
    modules = CourseModule.query.all()
    
    return render_template('textbook_search.html', 
                           sessions=sessions, 
                           classes=classes, 
                           modules=modules)

@textbook_bp.route('/textbook/<int:id>/delete', methods=['POST'])
def delete_session(id):
    session = TextbookSession.query.get_or_404(id)
    try:
        db.session.delete(session)
        db.session.commit()
        flash('Séance supprimée avec succès.', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erreur lors de la suppression: {str(e)}', 'error')
    return redirect(url_for('textbook.index'))

@textbook_bp.route('/textbook/delete-multiple', methods=['POST'])
def delete_multiple_sessions():
    session_ids = request.form.getlist('session_ids')
    if not session_ids:
        flash('Aucune séance sélectionnée.', 'warning')
        return redirect(url_for('textbook.index'))
    
    try:
        TextbookSession.query.filter(TextbookSession.id.in_(session_ids)).delete(synchronize_session=False)
        db.session.commit()
        flash(f'{len(session_ids)} séance(s) supprimée(s) avec succès.', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erreur lors de la suppression multiple: {str(e)}', 'error')

    # Redirection vers la liste du cahier de texte (avec filtre de classe s'il y a lieu)
    class_id = request.form.get('class_id') or request.args.get('class_id')
    if class_id:
        return redirect(url_for('textbook.index', class_id=class_id))
    return redirect(url_for('textbook.index'))

@textbook_bp.route('/textbook/class/<int:class_id>/print_full')
def print_full(class_id):
    """Génération de la vue officielle imprimable du Cahier de Texte complet de la classe."""
    classe = Class.query.get_or_404(class_id)
    institution = Institution.query.first() or Institution(name="Lycée d'Excellence")
    current_year = classe.school_year or get_current_school_year()

    sessions = TextbookSession.query.filter_by(
        class_id=classe.id
    ).order_by(TextbookSession.date.asc(), TextbookSession.id.asc()).all()

    # Enseignant principal déduit des séances ou valeur par défaut
    teacher_names = list(filter(None, [s.teacher_name for s in sessions]))
    main_teacher = teacher_names[0] if teacher_names else session.get('user_full_name', 'Enseignant')

    return render_template(
        'textbook_full_print.html',
        classe=classe,
        institution=institution,
        school_year=current_year,
        sessions=sessions,
        teacher_name=main_teacher,
        now=datetime.now()
    )

@textbook_bp.route('/textbook/export_excel')
def export_excel():
    """Exportation du Cahier de Texte au format Excel (.xlsx) avec openpyxl."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from werkzeug.utils import secure_filename

    year_id = request.args.get('year_id')
    class_id = request.args.get('class_id')
    search_q = request.args.get('q', '').strip()

    query = TextbookSession.query

    current_year = None
    if year_id:
        query = query.filter_by(school_year_id=year_id)
        current_year = db.session.get(SchoolYear, year_id)
    else:
        current_year = get_current_school_year()
        if current_year:
            query = query.filter_by(school_year_id=current_year.id)

    selected_class = None
    if class_id:
        query = query.filter_by(class_id=class_id)
        selected_class = db.session.get(Class, class_id)

    if search_q:
        search_term = f"%{search_q}%"
        query = query.join(CourseModule).join(CourseSection).filter(
            (TextbookSession.course_titles.ilike(search_term)) |
            (CourseModule.name.ilike(search_term)) |
            (CourseSection.name.ilike(search_term)) |
            (TextbookSession.teacher_name.ilike(search_term)) |
            (TextbookSession.remark.ilike(search_term))
        )

    sessions = query.order_by(TextbookSession.date.asc(), TextbookSession.id.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("Cahier de Texte")
    ws.title = "Cahier de Texte"
    ws.views.sheetView[0].showGridLines = True

    # Styles
    title_font = Font(name="Calibri", size=15, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="475569")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    data_font = Font(name="Calibri", size=10)
    bold_data_font = Font(name="Calibri", size=10, bold=True)
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Titre & Métadonnées
    classe_str = f"Classe : {selected_class.name}" if selected_class else "Toutes les classes"
    year_str = f"Année : {current_year.name}" if current_year else ""
    ws.merge_cells('A1:G1')
    ws['A1'] = "CAHIER DE TEXTE NUMÉRIQUE"
    ws['A1'].font = title_font
    ws['A1'].alignment = center_align

    ws.merge_cells('A2:G2')
    ws['A2'] = f"{classe_str} | {year_str} | Exporté le {datetime.now().strftime('%d/%m/%Y à %H:%M')}"
    ws['A2'].font = subtitle_font
    ws['A2'].alignment = center_align

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 8

    # En-têtes
    headers = ["#", "Date", "Classe", "Module & Chapitre", "Contenu / Titres du cours", "Objectifs & Compétences", "Enseignant"]
    header_row = 4
    ws.row_dimensions[header_row].height = 24

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Lignes
    for r_idx, s in enumerate(sessions, 1):
        curr_row = header_row + r_idx
        ws.row_dimensions[curr_row].height = 32
        is_zebra = (r_idx % 2 == 0)

        mod_sec = f"{s.module.name} / {s.section.name}" if s.module and s.section else "-"
        pedago = f"{len(s.objectives)} obj. / {len(s.competencies)} comp."

        row_vals = [
            r_idx,
            s.date.strftime('%d/%m/%Y'),
            s.classe.name if s.classe else "-",
            mod_sec,
            s.course_titles or "-",
            pedago,
            s.teacher_name or "-"
        ]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=curr_row, column=col_idx, value=val)
            cell.font = bold_data_font if col_idx in (2, 3) else data_font
            cell.alignment = left_align if col_idx in (4, 5) else center_align
            cell.border = thin_border
            if is_zebra:
                cell.fill = zebra_fill

    # Largeurs de colonnes
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 13
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 28
    ws.column_dimensions['E'].width = 38
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    class_part = secure_filename(selected_class.name) if selected_class else "Global"
    filename = f"Cahier_de_Texte_{class_part}_{datetime.now().strftime('%Y%m%d')}.xlsx"

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )

def generate_pdf(html_content, filename):
    if not weasyprint_available:
        # Fallback to HTML print view
        # Inject a print button and some basic styling if not present
        print_script = """
        <script>
            window.onload = function() {
                // Auto-print or just show the button
                // window.print(); 
            }
        </script>
        <style>
            .print-btn {
                position: fixed;
                top: 20px;
                right: 20px;
                padding: 10px 20px;
                background: #4e73df;
                color: white;
                border: none;
                border-radius: 5px;
                cursor: pointer;
                font-family: sans-serif;
                z-index: 1000;
            }
            @media print {
                .print-btn { display: none; }
            }
        </style>
        <button class="print-btn" onclick="window.print()">Imprimer / Enregistrer PDF</button>
        """
        if "</body>" in html_content:
            html_content = html_content.replace("</body>", f"{print_script}</body>")
        else:
            html_content += print_script
            
        return html_content
        
    try:
        if HTML is None:
            raise RuntimeError("WeasyPrint non disponible")
        html = HTML(string=html_content)
        pdf_bytes = html.write_pdf()
        
        response = make_response(pdf_bytes)
        response.headers['Content-Type'] = 'application/pdf'
        response.headers['Content-Disposition'] = f'inline; filename={filename}'
        return response
    except Exception as e:
        flash(f"Erreur lors de la génération PDF (WeasyPrint): {str(e)}", "error")
        return redirect(url_for('textbook.index'))
