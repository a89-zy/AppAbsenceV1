import io
import json
import math
import openpyxl
from datetime import datetime
from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, send_file
from database import (
    db, Class, Student, SchoolYear,
    SemesterActivityEvaluation, SemesterExamGrade,
    StudentBadge, AppSetting
)
from blueprints.activity import (
    get_current_school_year,
    get_semester_dates,
    get_activity_config,
    compute_student_activity_evaluation
)

evaluation_bp = Blueprint('evaluation', __name__)

DEFAULT_EVALUATION_WEIGHTS = {
    'cc1': 1.0,
    'cc2': 1.0,
    'cc3': 1.0,
    'cc4': 1.0,
    'activity': 1.0,
    'rounding': '0.01'
}

def get_evaluation_weights(class_id=None):
    """Récupère les coefficients d'évaluation (spécifiques à la classe ou généraux)."""
    if class_id:
        raw_class = AppSetting.get_value(f'eval_weights_class_{class_id}', '')
        if raw_class:
            try:
                data = json.loads(raw_class)
                weights = dict(DEFAULT_EVALUATION_WEIGHTS)
                weights.update(data)
                return weights
            except Exception:
                pass

    raw_global = AppSetting.get_value('eval_weights_global', '')
    if raw_global:
        try:
            data = json.loads(raw_global)
            weights = dict(DEFAULT_EVALUATION_WEIGHTS)
            weights.update(data)
            return weights
        except Exception:
            pass

    return dict(DEFAULT_EVALUATION_WEIGHTS)

def compute_weighted_subject_average(grades_dict, weights=None):
    """
    Calcule la moyenne semestrielle pondérée de la matière sur 20.
    grades_dict: {'cc1': float|None, 'cc2': float|None, 'cc3': float|None, 'cc4': float|None, 'activity': float|None}
    weights: {'cc1': float, 'cc2': float, 'cc3': float, 'cc4': float, 'activity': float, 'rounding': '0.01'|'0.25'|'0.5'}
    """
    if weights is None:
        weights = DEFAULT_EVALUATION_WEIGHTS

    total_points = 0.0
    total_weights = 0.0

    for key in ['cc1', 'cc2', 'cc3', 'cc4', 'activity']:
        val = grades_dict.get(key)
        if val is not None:
            w = float(weights.get(key, 1.0))
            if w > 0:
                total_points += val * w
                total_weights += w

    if total_weights <= 0:
        return None

    raw_avg = total_points / total_weights

    rounding_mode = weights.get('rounding', '0.01')
    if rounding_mode == '0.25':
        return round(round(raw_avg * 4) / 4, 2)
    elif rounding_mode == '0.5':
        return round(round(raw_avg * 2) / 2, 1)
    else:
        return round(raw_avg, 2)


@evaluation_bp.route('/')
def index():
    """Tableau de bord central du Module Évaluation : synthèse par classe et par semestre."""
    current_year = get_current_school_year()
    classes = Class.query.filter_by(school_year_id=current_year.id).order_by(Class.name).all() if current_year else Class.query.order_by(Class.name).all()

    selected_semester = request.args.get('semester', default=1, type=int)

    classes_summary = []
    total_students_count = 0
    total_grades_entered = 0

    if current_year:
        for c in classes:
            students = Student.query.filter_by(class_id=c.id).all()
            st_count = len(students)
            total_students_count += st_count

            # Exam grades for this class & semester
            exam_grades = SemesterExamGrade.query.filter_by(
                class_id=c.id,
                semester=selected_semester,
                school_year_id=current_year.id
            ).all()
            exam_map = {eg.student_id: eg for eg in exam_grades}

            # Activity evaluations
            act_evals = SemesterActivityEvaluation.query.filter_by(
                class_id=c.id,
                semester=selected_semester,
                school_year_id=current_year.id
            ).all()
            act_map = {ae.student_id: ae for ae in act_evals}

            cc1_done = sum(1 for s in students if exam_map.get(s.id) and exam_map[s.id].cc1 is not None)
            cc2_done = sum(1 for s in students if exam_map.get(s.id) and exam_map[s.id].cc2 is not None)
            cc3_done = sum(1 for s in students if exam_map.get(s.id) and exam_map[s.id].cc3 is not None)
            cc4_done = sum(1 for s in students if exam_map.get(s.id) and exam_map[s.id].cc4 is not None)
            act_done = sum(1 for s in students if act_map.get(s.id) and act_map[s.id].is_locked)

            total_grades_entered += (cc1_done + cc2_done + cc3_done + cc4_done + act_done)

            pct_progress = round(((cc1_done + cc2_done + cc3_done + cc4_done + act_done) / (st_count * 5) * 100), 1) if st_count > 0 else 0

            # Calcul moyenne générale de classe si notes présentes
            c_weights = get_evaluation_weights(c.id)
            averages = []
            for s in students:
                eg = exam_map.get(s.id)
                ae = act_map.get(s.id)
                g_dict = {
                    'cc1': eg.cc1 if eg else None,
                    'cc2': eg.cc2 if eg else None,
                    'cc3': eg.cc3 if eg else None,
                    'cc4': eg.cc4 if eg else None,
                    'activity': ae.final_score if ae and ae.final_score is not None else None
                }
                s_avg = compute_weighted_subject_average(g_dict, c_weights)
                if s_avg is not None:
                    averages.append(s_avg)

            class_avg = round(sum(averages) / len(averages), 2) if averages else None

            classes_summary.append({
                'class': c,
                'student_count': st_count,
                'cc1_done': cc1_done,
                'cc2_done': cc2_done,
                'cc3_done': cc3_done,
                'cc4_done': cc4_done,
                'act_done': act_done,
                'pct_progress': pct_progress,
                'class_avg': class_avg
            })

    return render_template(
        'evaluation/index.html',
        classes=classes,
        classes_summary=classes_summary,
        selected_semester=selected_semester,
        total_students_count=total_students_count,
        total_grades_entered=total_grades_entered,
        current_year=current_year
    )


@evaluation_bp.route('/grades', methods=['GET', 'POST'])
@evaluation_bp.route('/class/<int:class_id>', methods=['GET', 'POST'])
def grades_grid(class_id=None):
    """Carnet des 5 Notes de l'Informatique : CC1, CC2, CC3, CC4, Note d'Activité et Moyenne Générale."""
    current_year = get_current_school_year()
    classes = Class.query.filter_by(school_year_id=current_year.id).order_by(Class.name).all() if current_year else Class.query.order_by(Class.name).all()

    req_class_id = class_id or request.args.get('class_id', type=int)
    semester = request.args.get('semester', default=1, type=int)

    selected_class = None
    if req_class_id:
        selected_class = Class.query.get(req_class_id)
    if not selected_class and classes:
        selected_class = classes[0]

    cfg = get_activity_config()
    start_date, end_date = get_semester_dates(semester, current_year)

    evaluations = []
    class_stats = {
        'count': 0,
        'avg': 0.0,
        'min': 0.0,
        'max': 0.0,
        'pass_count': 0,
        'pass_rate': 0.0,
        'excellent': 0,
        'good': 0,
        'fair': 0,
        'insufficient': 0
    }

    if request.method == 'POST':
        if selected_class and current_year:
            students = Student.query.filter_by(class_id=selected_class.id).order_by(*Student.default_order()).all()
            for st in students:
                # 1. Traitement Note d'activité
                manual_adj_str = request.form.get(f'manual_adj_{st.id}', '0.0')
                appr = request.form.get(f'appreciation_{st.id}', '').strip()
                try:
                    manual_adj = float(manual_adj_str)
                except ValueError:
                    manual_adj = 0.0

                calc = compute_student_activity_evaluation(st, semester, selected_class, current_year)
                final_score = round(max(0.0, min(20.0, calc['base_score'] + calc['bonus_score'] - calc['malus_score'] + manual_adj)), 2)

                saved_eval = SemesterActivityEvaluation.query.filter_by(
                    student_id=st.id,
                    class_id=selected_class.id,
                    semester=semester,
                    school_year_id=current_year.id
                ).first()

                if not saved_eval:
                    saved_eval = SemesterActivityEvaluation(
                        student_id=st.id,
                        class_id=selected_class.id,
                        semester=semester,
                        school_year_id=current_year.id,
                        base_score=calc['base_score'],
                        bonus_score=calc['bonus_score'],
                        malus_score=calc['malus_score'],
                        manual_adjustment=manual_adj,
                        final_score=final_score,
                        appreciation=appr or calc['appreciation'],
                        is_locked=True
                    )
                    db.session.add(saved_eval)
                else:
                    saved_eval.base_score = calc['base_score']
                    saved_eval.bonus_score = calc['bonus_score']
                    saved_eval.malus_score = calc['malus_score']
                    saved_eval.manual_adjustment = manual_adj
                    saved_eval.final_score = final_score
                    saved_eval.appreciation = appr or calc['appreciation']
                    saved_eval.is_locked = True

                # 2. Traitement des 4 Contrôles Continus (CC1 à CC4)
                def parse_grade(val_str):
                    if not val_str or not val_str.strip():
                        return None
                    try:
                        v = float(val_str.replace(',', '.'))
                        return round(max(0.0, min(20.0, v)), 2)
                    except ValueError:
                        return None

                cc1_val = parse_grade(request.form.get(f'cc1_{st.id}'))
                cc2_val = parse_grade(request.form.get(f'cc2_{st.id}'))
                cc3_val = parse_grade(request.form.get(f'cc3_{st.id}'))
                cc4_val = parse_grade(request.form.get(f'cc4_{st.id}'))

                exam_rec = SemesterExamGrade.query.filter_by(
                    student_id=st.id,
                    semester=semester,
                    school_year_id=current_year.id
                ).first()

                if not exam_rec:
                    exam_rec = SemesterExamGrade(
                        student_id=st.id,
                        class_id=selected_class.id,
                        school_year_id=current_year.id,
                        semester=semester,
                        cc1=cc1_val,
                        cc2=cc2_val,
                        cc3=cc3_val,
                        cc4=cc4_val
                    )
                    db.session.add(exam_rec)
                else:
                    exam_rec.cc1 = cc1_val
                    exam_rec.cc2 = cc2_val
                    exam_rec.cc3 = cc3_val
                    exam_rec.cc4 = cc4_val

            db.session.commit()
            flash("Carnet de notes semestriel enregistré : Contrôles Continus (CC1-CC4) et Note d'Activité validés avec succès.", "success")
            return redirect(url_for('evaluation.grades_grid', class_id=selected_class.id, semester=semester))

    # Chargement en mode GET
    if selected_class and current_year:
        students = Student.query.filter_by(class_id=selected_class.id).order_by(*Student.default_order()).all()
        saved_map = {
            e.student_id: e 
            for e in SemesterActivityEvaluation.query.filter_by(
                class_id=selected_class.id, 
                semester=semester, 
                school_year_id=current_year.id
            ).all()
        }
        exam_map = {
            e.student_id: e 
            for e in SemesterExamGrade.query.filter_by(
                class_id=selected_class.id,
                semester=semester,
                school_year_id=current_year.id
            ).all()
        }

        weights = get_evaluation_weights(selected_class.id if selected_class else None)
        scores_list = []
        for st in students:
            eval_data = compute_student_activity_evaluation(st, semester, selected_class, current_year, saved_map.get(st.id))
            ex_rec = exam_map.get(st.id)
            eval_data['cc1'] = ex_rec.cc1 if ex_rec else None
            eval_data['cc2'] = ex_rec.cc2 if ex_rec else None
            eval_data['cc3'] = ex_rec.cc3 if ex_rec else None
            eval_data['cc4'] = ex_rec.cc4 if ex_rec else None
            eval_data['cc_avg'] = ex_rec.cc_average if ex_rec else None

            # Moyenne générale des 5 notes de la matière pondérée
            g_dict = {
                'cc1': eval_data['cc1'],
                'cc2': eval_data['cc2'],
                'cc3': eval_data['cc3'],
                'cc4': eval_data['cc4'],
                'activity': eval_data['final_score']
            }
            w_avg = compute_weighted_subject_average(g_dict, weights)
            eval_data['subject_average'] = w_avg if w_avg is not None else (eval_data['final_score'] or 0.0)

            evaluations.append(eval_data)
            scores_list.append(eval_data['subject_average'])

            sa = eval_data['subject_average']
            if sa >= 16: class_stats['excellent'] += 1
            elif sa >= 12: class_stats['good'] += 1
            elif sa >= 10: class_stats['fair'] += 1
            else: class_stats['insufficient'] += 1

        if scores_list:
            class_stats['count'] = len(scores_list)
            class_stats['avg'] = round(sum(scores_list) / len(scores_list), 2)
            class_stats['min'] = min(scores_list)
            class_stats['max'] = max(scores_list)
            class_stats['pass_count'] = sum(1 for s in scores_list if s >= 10.0)
            class_stats['pass_rate'] = round((class_stats['pass_count'] / len(scores_list)) * 100.0, 1)

    return render_template(
        'evaluation/grades_grid.html',
        classes=classes,
        selected_class=selected_class,
        semester=semester,
        start_date=start_date,
        end_date=end_date,
        cfg=cfg,
        evaluations=evaluations,
        stats=class_stats,
        current_year=current_year,
        weights=weights
    )


@evaluation_bp.route('/api/save-exam-grades', methods=['POST'])
def api_save_exam_grades():
    """Sauvegarde rapide en AJAX des notes de contrôles continus (CC1 à CC4) d'un étudiant."""
    data = request.get_json() or {}
    student_id = data.get('student_id')
    semester = int(data.get('semester', 1))
    class_id = data.get('class_id')

    if not student_id:
        return jsonify({'success': False, 'error': 'ID étudiant manquant'}), 400

    current_year = get_current_school_year()
    if not current_year:
        return jsonify({'success': False, 'error': 'Aucune année scolaire active'}), 400

    def clean_val(v):
        if v is None or v == '':
            return None
        try:
            return round(max(0.0, min(20.0, float(str(v).replace(',', '.')))), 2)
        except (ValueError, TypeError):
            return None

    cc1 = clean_val(data.get('cc1'))
    cc2 = clean_val(data.get('cc2'))
    cc3 = clean_val(data.get('cc3'))
    cc4 = clean_val(data.get('cc4'))

    rec = SemesterExamGrade.query.filter_by(
        student_id=student_id,
        semester=semester,
        school_year_id=current_year.id
    ).first()

    if not rec:
        st = Student.query.get(student_id)
        rec = SemesterExamGrade(
            student_id=student_id,
            class_id=class_id or (st.class_id if st else 1),
            school_year_id=current_year.id,
            semester=semester,
            cc1=cc1,
            cc2=cc2,
            cc3=cc3,
            cc4=cc4
        )
        db.session.add(rec)
    else:
        rec.cc1 = cc1
        rec.cc2 = cc2
        rec.cc3 = cc3
        rec.cc4 = cc4

    db.session.commit()
    return jsonify({
        'success': True,
        'message': 'Notes enregistrées',
        'cc_average': rec.cc_average
    })


@evaluation_bp.route('/massar', methods=['GET', 'POST'])
def massar_hub():
    """Hub dédié à la Passerelle MASSAR : injection automatique des 5 notes dans le fichier Excel officiel."""
    current_year = get_current_school_year()
    classes = Class.query.filter_by(school_year_id=current_year.id).order_by(Class.name).all() if current_year else Class.query.order_by(Class.name).all()

    class_id = request.form.get('class_id') or request.args.get('class_id', type=int)
    semester = int(request.form.get('semester') or request.args.get('semester', 1))

    selected_class = None
    if class_id:
        selected_class = Class.query.get(int(class_id))
    if not selected_class and classes:
        selected_class = classes[0]

    # Traitement POST d'injection
    if request.method == 'POST' and 'massar_file' in request.files:
        file = request.files['massar_file']
        if not file or file.filename == '':
            flash('Veuillez sélectionner un fichier Excel (.xlsx) téléchargé depuis MASSAR.', 'danger')
            return redirect(url_for('evaluation.massar_hub', class_id=selected_class.id if selected_class else None, semester=semester))

        if not selected_class or not current_year:
            flash('Classe ou année scolaire invalide.', 'danger')
            return redirect(url_for('evaluation.massar_hub'))

        try:
            wb = openpyxl.load_workbook(file)
            ws = wb.active

            # 1. Détection des colonnes officielles dans la feuille MASSAR
            cne_col = None
            cc1_col = None
            cc2_col = None
            cc3_col = None
            cc4_col = None
            act_col = None
            header_row_idx = 1

            for r in range(1, min(10, ws.max_row + 1)):
                for c in range(1, ws.max_column + 1):
                    val = str(ws.cell(row=r, column=c).value or '').strip()
                    val_lower = val.lower()

                    if 'الرمز' in val or 'cne' in val_lower or 'massar' in val_lower or 'code' in val_lower:
                        cne_col = c
                        header_row_idx = r
                    elif 'المراقبة 1' in val or 'cc1' in val_lower or 'controle 1' in val_lower:
                        cc1_col = c
                    elif 'المراقبة 2' in val or 'cc2' in val_lower or 'controle 2' in val_lower:
                        cc2_col = c
                    elif 'المراقبة 3' in val or 'cc3' in val_lower or 'controle 3' in val_lower:
                        cc3_col = c
                    elif 'المراقبة 4' in val or 'cc4' in val_lower or 'controle 4' in val_lower:
                        cc4_col = c
                    elif 'الأنشطة المندمجة' in val or 'الأنشطة' in val or 'activite' in val_lower:
                        act_col = c

            if not cne_col:
                cne_col = 2 # Défaut standard MASSAR

            # Données de la classe dans l'application
            students = Student.query.filter_by(class_id=selected_class.id).all()
            st_by_cne = {s.cne.strip().upper(): s for s in students if s.cne}

            exam_grades = SemesterExamGrade.query.filter_by(
                class_id=selected_class.id,
                semester=semester,
                school_year_id=current_year.id
            ).all()
            exam_map = {eg.student_id: eg for eg in exam_grades}

            act_evals = SemesterActivityEvaluation.query.filter_by(
                class_id=selected_class.id,
                semester=semester,
                school_year_id=current_year.id
            ).all()
            act_map = {ae.student_id: ae for ae in act_evals}

            injected_count = 0

            # 2. Injection ligne par ligne
            for row in range(header_row_idx + 1, ws.max_row + 1):
                cell_cne = ws.cell(row=row, column=cne_col).value
                if not cell_cne:
                    continue
                code_cne = str(cell_cne).strip().upper()

                if code_cne in st_by_cne:
                    st = st_by_cne[code_cne]
                    ex = exam_map.get(st.id)
                    act = act_map.get(st.id)

                    # Note d'Activité
                    if act_col:
                        act_score = act.final_score if (act and act.final_score is not None) else None
                        if act_score is not None:
                            ws.cell(row=row, column=act_col, value=act_score)

                    # Contrôles Continus (CC1 à CC4)
                    if ex:
                        if cc1_col and ex.cc1 is not None:
                            ws.cell(row=row, column=cc1_col, value=ex.cc1)
                        if cc2_col and ex.cc2 is not None:
                            ws.cell(row=row, column=cc2_col, value=ex.cc2)
                        if cc3_col and ex.cc3 is not None:
                            ws.cell(row=row, column=cc3_col, value=ex.cc3)
                        if cc4_col and ex.cc4 is not None:
                            ws.cell(row=row, column=cc4_col, value=ex.cc4)

                    injected_count += 1

            # Sauvegarde en mémoire et renvoi du fichier
            out_stream = io.BytesIO()
            wb.save(out_stream)
            out_stream.seek(0)

            clean_filename = f"MASSAR_Injecte_{selected_class.name}_S{semester}.xlsx"
            return send_file(
                out_stream,
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                as_attachment=True,
                download_name=clean_filename
            )

        except Exception as e:
            flash(f"Erreur lors du traitement du fichier MASSAR : {str(e)}", 'danger')
            return redirect(url_for('evaluation.massar_hub', class_id=selected_class.id if selected_class else None, semester=semester))

    return render_template(
        'evaluation/massar.html',
        classes=classes,
        selected_class=selected_class,
        semester=semester,
        current_year=current_year
    )


@evaluation_bp.route('/quick-entry', methods=['GET', 'POST'])
def quick_entry():
    """Interface focalisée de saisie ultra-rapide pour un seul Contrôle Continu (CC1, CC2, CC3 ou CC4)."""
    current_year = get_current_school_year()
    classes = Class.query.filter_by(school_year_id=current_year.id).order_by(Class.name).all() if current_year else Class.query.order_by(Class.name).all()

    class_id = request.form.get('class_id') or request.args.get('class_id', type=int)
    semester = int(request.form.get('semester') or request.args.get('semester', 1))
    cc_number = int(request.form.get('cc_number') or request.args.get('cc_number', 1))
    if cc_number not in [1, 2, 3, 4]:
        cc_number = 1

    selected_class = None
    if class_id:
        selected_class = Class.query.get(int(class_id))
    if not selected_class and classes:
        selected_class = classes[0]

    if request.method == 'POST':
        if selected_class and current_year:
            students = Student.query.filter_by(class_id=selected_class.id).order_by(*Student.default_order()).all()
            for st in students:
                grade_str = request.form.get(f'grade_{st.id}', '').strip()
                is_absent = request.form.get(f'absent_{st.id}') == '1'

                val = None
                if not is_absent and grade_str != '':
                    try:
                        val = round(max(0.0, min(20.0, float(grade_str.replace(',', '.')))), 2)
                    except ValueError:
                        val = None

                rec = SemesterExamGrade.query.filter_by(
                    student_id=st.id,
                    semester=semester,
                    school_year_id=current_year.id
                ).first()

                if not rec:
                    rec = SemesterExamGrade(
                        student_id=st.id,
                        class_id=selected_class.id,
                        school_year_id=current_year.id,
                        semester=semester
                    )
                    db.session.add(rec)

                if cc_number == 1: rec.cc1 = val
                elif cc_number == 2: rec.cc2 = val
                elif cc_number == 3: rec.cc3 = val
                elif cc_number == 4: rec.cc4 = val

            db.session.commit()
            flash(f"Notes du Contrôle Continu N°{cc_number} enregistrées avec succès pour la classe {selected_class.name}.", "success")
            return redirect(url_for('evaluation.quick_entry', class_id=selected_class.id, semester=semester, cc_number=cc_number))

    # Mode GET
    student_items = []
    stats = {
        'count': 0,
        'total_students': 0,
        'avg': 0.0,
        'min': 0.0,
        'max': 0.0,
        'pass_count': 0,
        'pass_rate': 0.0
    }

    if selected_class and current_year:
        students = Student.query.filter_by(class_id=selected_class.id).order_by(*Student.default_order()).all()
        stats['total_students'] = len(students)

        exam_map = {
            eg.student_id: eg 
            for eg in SemesterExamGrade.query.filter_by(
                class_id=selected_class.id,
                semester=semester,
                school_year_id=current_year.id
            ).all()
        }

        grades_list = []
        for st in students:
            eg = exam_map.get(st.id)
            val = getattr(eg, f'cc{cc_number}', None) if eg else None
            if val is not None:
                grades_list.append(val)

            student_items.append({
                'student': st,
                'grade': val,
                'has_grade': val is not None
            })

        if grades_list:
            stats['count'] = len(grades_list)
            stats['avg'] = round(sum(grades_list) / len(grades_list), 2)
            stats['min'] = min(grades_list)
            stats['max'] = max(grades_list)
            stats['pass_count'] = sum(1 for g in grades_list if g >= 10.0)
            stats['pass_rate'] = round((stats['pass_count'] / len(grades_list)) * 100.0, 1)

    return render_template(
        'evaluation/quick_entry.html',
        classes=classes,
        selected_class=selected_class,
        semester=semester,
        cc_number=cc_number,
        student_items=student_items,
        stats=stats,
        current_year=current_year
    )


@evaluation_bp.route('/api/save-single-cc', methods=['POST'])
def api_save_single_cc():
    """Sauvegarde asynchrone AJAX lors de la saisie d'une note de CC, avec recalcul live des statistiques de la classe."""
    data = request.get_json() or {}
    student_id = data.get('student_id')
    class_id = data.get('class_id')
    semester = int(data.get('semester', 1))
    cc_number = int(data.get('cc_number', 1))
    is_absent = bool(data.get('is_absent', False))
    raw_val = data.get('grade')

    if not student_id:
        return jsonify({'success': False, 'error': 'ID étudiant manquant'}), 400

    current_year = get_current_school_year()
    if not current_year:
        return jsonify({'success': False, 'error': 'Aucune année scolaire active'}), 400

    val = None
    if not is_absent and raw_val is not None and str(raw_val).strip() != '':
        try:
            val = round(max(0.0, min(20.0, float(str(raw_val).replace(',', '.')))), 2)
        except (ValueError, TypeError):
            val = None

    rec = SemesterExamGrade.query.filter_by(
        student_id=student_id,
        semester=semester,
        school_year_id=current_year.id
    ).first()

    if not rec:
        st = Student.query.get(student_id)
        rec = SemesterExamGrade(
            student_id=student_id,
            class_id=class_id or (st.class_id if st else 1),
            school_year_id=current_year.id,
            semester=semester
        )
        db.session.add(rec)

    setattr(rec, f'cc{cc_number}', val)
    db.session.commit()

    # Recalcul des statistiques en direct pour la classe
    class_obj_id = class_id or rec.class_id
    all_exams = SemesterExamGrade.query.filter_by(
        class_id=class_obj_id,
        semester=semester,
        school_year_id=current_year.id
    ).all()

    grades_list = []
    for eg in all_exams:
        g = getattr(eg, f'cc{cc_number}', None)
        if g is not None:
            grades_list.append(g)

    total_students = Student.query.filter_by(class_id=class_obj_id).count()

    stats = {
        'count': len(grades_list),
        'total_students': total_students,
        'avg': round(sum(grades_list) / len(grades_list), 2) if grades_list else 0.0,
        'min': min(grades_list) if grades_list else 0.0,
        'max': max(grades_list) if grades_list else 0.0,
        'pass_count': sum(1 for g in grades_list if g >= 10.0),
        'pass_rate': round((sum(1 for g in grades_list if g >= 10.0) / len(grades_list)) * 100.0, 1) if grades_list else 0.0
    }

    return jsonify({
        'success': True,
        'grade': val,
        'cc_average': rec.cc_average,
        'stats': stats
    })


@evaluation_bp.route('/analytics')
def analytics():
    """Tableau de Bord Statistique & Détection Automatique de Remédiation."""
    current_year = get_current_school_year()
    classes = Class.query.filter_by(school_year_id=current_year.id).order_by(Class.name).all() if current_year else Class.query.order_by(Class.name).all()

    class_id = request.args.get('class_id', type=int)
    semester = request.args.get('semester', default=1, type=int)

    selected_class = None
    if class_id:
        selected_class = Class.query.get(class_id)
    if not selected_class and classes:
        selected_class = classes[0]

    class_weights = get_evaluation_weights(selected_class.id if selected_class else None)

    class_stats = {
        'total_students': 0,
        'evaluated_count': 0,
        'subject_avg': 0.0,
        'min_score': 0.0,
        'max_score': 0.0,
        'std_dev': 0.0,
        'pass_count': 0,
        'pass_rate': 0.0,
        'remediation_count': 0
    }

    # Distributions par niveau de maîtrise
    distribution = {
        'excellent': {'count': 0, 'pct': 0.0, 'label': 'Très Bien (≥ 16)', 'color': 'success'},
        'good': {'count': 0, 'pct': 0.0, 'label': 'Bien / Assez Bien (12 - 16)', 'color': 'primary'},
        'fair': {'count': 0, 'pct': 0.0, 'label': 'Passable / Fragile (10 - 12)', 'color': 'info'},
        'warning': {'count': 0, 'pct': 0.0, 'label': 'Vigilance (8 - 10)', 'color': 'warning'},
        'critical': {'count': 0, 'pct': 0.0, 'label': 'Difficulté Majeure (< 8)', 'color': 'danger'}
    }

    # Évolution comparative : CC1..CC4 & Activité
    exam_series = {
        'cc1': {'label': 'CC1', 'scores': [], 'avg': 0.0, 'pass_rate': 0.0},
        'cc2': {'label': 'CC2', 'scores': [], 'avg': 0.0, 'pass_rate': 0.0},
        'cc3': {'label': 'CC3', 'scores': [], 'avg': 0.0, 'pass_rate': 0.0},
        'cc4': {'label': 'CC4', 'scores': [], 'avg': 0.0, 'pass_rate': 0.0},
        'act': {'label': 'Activité', 'scores': [], 'avg': 0.0, 'pass_rate': 0.0}
    }

    remediation_list = []
    honor_roll = []

    if selected_class and current_year:
        students = Student.query.filter_by(class_id=selected_class.id).order_by(*Student.default_order()).all()
        class_stats['total_students'] = len(students)

        exam_map = {
            eg.student_id: eg 
            for eg in SemesterExamGrade.query.filter_by(
                class_id=selected_class.id,
                semester=semester,
                school_year_id=current_year.id
            ).all()
        }

        act_map = {
            ae.student_id: ae 
            for ae in SemesterActivityEvaluation.query.filter_by(
                class_id=selected_class.id,
                semester=semester,
                school_year_id=current_year.id
            ).all()
        }

        student_analyses = []
        all_subject_scores = []

        class_weights = get_evaluation_weights(selected_class.id if selected_class else None)
        for st in students:
            ex = exam_map.get(st.id)
            act = act_map.get(st.id)

            cc1 = ex.cc1 if ex else None
            cc2 = ex.cc2 if ex else None
            cc3 = ex.cc3 if ex else None
            cc4 = ex.cc4 if ex else None
            act_score = act.final_score if act and act.final_score is not None else 10.0

            if cc1 is not None: exam_series['cc1']['scores'].append(cc1)
            if cc2 is not None: exam_series['cc2']['scores'].append(cc2)
            if cc3 is not None: exam_series['cc3']['scores'].append(cc3)
            if cc4 is not None: exam_series['cc4']['scores'].append(cc4)
            if act_score is not None: exam_series['act']['scores'].append(act_score)

            g_dict = {
                'cc1': cc1,
                'cc2': cc2,
                'cc3': cc3,
                'cc4': cc4,
                'activity': act_score
            }
            subj_avg = compute_weighted_subject_average(g_dict, class_weights)

            if subj_avg is not None:
                all_subject_scores.append(subj_avg)

                # Répartition par tranches
                if subj_avg >= 16.0: distribution['excellent']['count'] += 1
                elif subj_avg >= 12.0: distribution['good']['count'] += 1
                elif subj_avg >= 10.0: distribution['fair']['count'] += 1
                elif subj_avg >= 8.0: distribution['warning']['count'] += 1
                else: distribution['critical']['count'] += 1

            failed_ccs = []
            if cc1 is not None and cc1 < 10.0: failed_ccs.append(('CC1', cc1))
            if cc2 is not None and cc2 < 10.0: failed_ccs.append(('CC2', cc2))
            if cc3 is not None and cc3 < 10.0: failed_ccs.append(('CC3', cc3))
            if cc4 is not None and cc4 < 10.0: failed_ccs.append(('CC4', cc4))

            st_data = {
                'student': st,
                'cc1': cc1,
                'cc2': cc2,
                'cc3': cc3,
                'cc4': cc4,
                'act_score': act_score,
                'subj_avg': subj_avg,
                'failed_ccs': failed_ccs,
                'malus_score': act.malus_score if act else 0.0
            }
            student_analyses.append(st_data)

            # Identification des élèves ayant besoin de remédiation
            if subj_avg is not None and (subj_avg < 10.0 or len(failed_ccs) >= 2):
                priority = 'critical' if (subj_avg < 8.0 or len(failed_ccs) >= 3) else 'warning'
                action = "Exercices guidés & révision ciblée"
                if act and act.malus_score >= 1.5:
                    action += " + Contrat d'assiduité"
                elif subj_avg < 8.0:
                    action = "Remédiation renforcée & tutorat par binôme"

                remediation_list.append({
                    'student': st,
                    'subj_avg': subj_avg,
                    'failed_ccs': failed_ccs,
                    'priority': priority,
                    'action': action
                })

        # Trier remédiation par moyenne croissante
        remediation_list.sort(key=lambda x: (x['subj_avg'] if x['subj_avg'] is not None else 0))

        # Tableau d'honneur (Top 5)
        honor_roll = [s for s in student_analyses if s['subj_avg'] is not None and s['subj_avg'] >= 10.0]
        honor_roll.sort(key=lambda x: x['subj_avg'], reverse=True)
        honor_roll = honor_roll[:5]

        # Statistiques générales de la classe
        if all_subject_scores:
            n = len(all_subject_scores)
            class_stats['evaluated_count'] = n
            mean = sum(all_subject_scores) / n
            class_stats['subject_avg'] = round(mean, 2)
            class_stats['min_score'] = min(all_subject_scores)
            class_stats['max_score'] = max(all_subject_scores)
            class_stats['pass_count'] = sum(1 for s in all_subject_scores if s >= 10.0)
            class_stats['pass_rate'] = round((class_stats['pass_count'] / n) * 100.0, 1)
            class_stats['remediation_count'] = len(remediation_list)

            # Écart-type (mesure de l'hétérogénéité de la classe)
            variance = sum((x - mean) ** 2 for x in all_subject_scores) / n
            class_stats['std_dev'] = round(math.sqrt(variance), 2)

            for k in distribution:
                distribution[k]['pct'] = round((distribution[k]['count'] / n) * 100.0, 1)

        # Calcul des moyennes des épreuves CC1..CC4
        for k in exam_series:
            scores = exam_series[k]['scores']
            if scores:
                exam_series[k]['avg'] = round(sum(scores) / len(scores), 2)
                pass_cnt = sum(1 for s in scores if s >= 10.0)
                exam_series[k]['pass_rate'] = round((pass_cnt / len(scores)) * 100.0, 1)

    return render_template(
        'evaluation/analytics.html',
        classes=classes,
        selected_class=selected_class,
        semester=semester,
        stats=class_stats,
        distribution=distribution,
        exam_series=exam_series,
        remediation_list=remediation_list,
        honor_roll=honor_roll,
        current_year=current_year,
        weights=class_weights
    )


def build_student_report_dict(student, semester, current_year, class_students_scores=None):
    """Construit les données complètes pour la fiche individuelle de notes d'un élève."""
    c = student.student_class
    exam_rec = SemesterExamGrade.query.filter_by(
        student_id=student.id,
        semester=semester,
        school_year_id=current_year.id
    ).first()

    act_rec = SemesterActivityEvaluation.query.filter_by(
        student_id=student.id,
        semester=semester,
        school_year_id=current_year.id
    ).first()

    cc1 = exam_rec.cc1 if exam_rec else None
    cc2 = exam_rec.cc2 if exam_rec else None
    cc3 = exam_rec.cc3 if exam_rec else None
    cc4 = exam_rec.cc4 if exam_rec else None

    # Calcul activité Formule B
    act_calc = compute_student_activity_evaluation(student, semester, c, current_year, act_rec)
    act_final = act_rec.final_score if act_rec and act_rec.final_score is not None else act_calc['final_score']

    weights = get_evaluation_weights(c.id if c else None)
    g_dict = {
        'cc1': cc1,
        'cc2': cc2,
        'cc3': cc3,
        'cc4': cc4,
        'activity': act_final
    }
    subject_avg = compute_weighted_subject_average(g_dict, weights)

    if subject_avg is not None:
        if subject_avg >= 16.0: mention = "Très Bien"
        elif subject_avg >= 14.0: mention = "Bien"
        elif subject_avg >= 12.0: mention = "Assez Bien"
        elif subject_avg >= 10.0: mention = "Passable"
        elif subject_avg >= 8.0: mention = "Insuffisant"
        else: mention = "Non Validé"
    else:
        mention = "Non Évalué"

    # Badges de l'élève
    badges = StudentBadge.query.filter_by(student_id=student.id).order_by(StudentBadge.awarded_at.desc()).all()

    appreciation = act_rec.appreciation if act_rec and act_rec.appreciation else act_calc.get('appreciation', '')

    # Calcul du rang et statistiques comparatives
    rank = None
    class_avg = None
    class_min = None
    class_max = None
    total_class = 0

    if class_students_scores:
        total_class = len(class_students_scores)
        sorted_scores = sorted(class_students_scores, key=lambda x: (x[1] if x[1] is not None else -1), reverse=True)
        for idx, (st_id, score) in enumerate(sorted_scores, 1):
            if st_id == student.id and score is not None:
                rank = idx
                break
        valid_class_scores = [s for _, s in class_students_scores if s is not None]
        if valid_class_scores:
            class_avg = round(sum(valid_class_scores) / len(valid_class_scores), 2)
            class_min = min(valid_class_scores)
            class_max = max(valid_class_scores)

    return {
        'student': student,
        'classe': c,
        'semester': semester,
        'current_year': current_year,
        'cc1': cc1,
        'cc2': cc2,
        'cc3': cc3,
        'cc4': cc4,
        'act_final': act_final,
        'act_calc': act_calc,
        'subject_avg': subject_avg,
        'mention': mention,
        'rank': rank,
        'total_class': total_class,
        'class_avg': class_avg,
        'class_min': class_min,
        'class_max': class_max,
        'badges': badges,
        'appreciation': appreciation,
        'weights': weights
    }


@evaluation_bp.route('/student/<int:student_id>/report')
def student_report(student_id):
    """Fiche individuelle de notes et bilan semestriel pour un élève."""
    current_year = get_current_school_year()
    student = Student.query.get_or_404(student_id)
    c = student.student_class
    semester = request.args.get('semester', default=1, type=int)

    all_students = Student.query.filter_by(class_id=c.id).all()
    c_weights = get_evaluation_weights(c.id if c else None)
    class_scores = []
    for s in all_students:
        ex = SemesterExamGrade.query.filter_by(student_id=s.id, semester=semester, school_year_id=current_year.id).first()
        ac = SemesterActivityEvaluation.query.filter_by(student_id=s.id, semester=semester, school_year_id=current_year.id).first()
        g_dict = {
            'cc1': ex.cc1 if ex else None,
            'cc2': ex.cc2 if ex else None,
            'cc3': ex.cc3 if ex else None,
            'cc4': ex.cc4 if ex else None,
            'activity': ac.final_score if ac and ac.final_score is not None else 10.0
        }
        s_avg = compute_weighted_subject_average(g_dict, c_weights)
        class_scores.append((s.id, s_avg))

    report_data = build_student_report_dict(student, semester, current_year, class_scores)

    return render_template(
        'evaluation/student_report.html',
        reports=[report_data],
        is_batch=False,
        current_year=current_year,
        semester=semester,
        selected_class=c,
        weights=c_weights
    )


@evaluation_bp.route('/class/<int:class_id>/reports')
def class_reports(class_id):
    """Génération groupée de toutes les fiches individuelles de notes de la classe (Impression par lot)."""
    current_year = get_current_school_year()
    c = Class.query.get_or_404(class_id)
    semester = request.args.get('semester', default=1, type=int)

    students = Student.query.filter_by(class_id=c.id).order_by(*Student.default_order()).all()
    c_weights = get_evaluation_weights(c.id if c else None)

    class_scores = []
    for s in students:
        ex = SemesterExamGrade.query.filter_by(student_id=s.id, semester=semester, school_year_id=current_year.id).first()
        ac = SemesterActivityEvaluation.query.filter_by(student_id=s.id, semester=semester, school_year_id=current_year.id).first()
        g_dict = {
            'cc1': ex.cc1 if ex else None,
            'cc2': ex.cc2 if ex else None,
            'cc3': ex.cc3 if ex else None,
            'cc4': ex.cc4 if ex else None,
            'activity': ac.final_score if ac and ac.final_score is not None else 10.0
        }
        s_avg = compute_weighted_subject_average(g_dict, c_weights)
        class_scores.append((s.id, s_avg))

    reports = [build_student_report_dict(s, semester, current_year, class_scores) for s in students]

    return render_template(
        'evaluation/student_report.html',
        reports=reports,
        is_batch=True,
        current_year=current_year,
        semester=semester,
        selected_class=c
    )


@evaluation_bp.route('/api/save-appreciation', methods=['POST'])
def api_save_appreciation():
    """Sauvegarde rapide d'une appréciation individualisée."""
    data = request.get_json() or {}
    student_id = data.get('student_id')
    semester = int(data.get('semester', 1))
    appreciation = data.get('appreciation', '').strip()

    if not student_id:
        return jsonify({'success': False, 'error': 'ID étudiant manquant'}), 400

    current_year = get_current_school_year()
    if not current_year:
        return jsonify({'success': False, 'error': 'Aucune année active'}), 400

    st = Student.query.get_or_404(student_id)
    rec = SemesterActivityEvaluation.query.filter_by(
        student_id=student_id,
        semester=semester,
        school_year_id=current_year.id
    ).first()

    if not rec:
        calc = compute_student_activity_evaluation(st, semester, st.student_class, current_year)
        rec = SemesterActivityEvaluation(
            student_id=student_id,
            class_id=st.class_id,
            school_year_id=current_year.id,
            semester=semester,
            base_score=calc['base_score'],
            bonus_score=calc['bonus_score'],
            malus_score=calc['malus_score'],
            final_score=calc['final_score'],
            appreciation=appreciation,
            is_locked=True
        )
        db.session.add(rec)
    else:
        rec.appreciation = appreciation

    db.session.commit()
    return jsonify({'success': True, 'message': 'Appréciation enregistrée avec succès.'})


@evaluation_bp.route('/make-up')
def makeup_manager():
    """Gestionnaire des contrôles continus à rattraper (élèves sans note ou absents)."""
    current_year = get_current_school_year()
    classes = Class.query.filter_by(school_year_id=current_year.id).order_by(Class.name).all() if current_year else Class.query.order_by(Class.name).all()

    class_id = request.args.get('class_id', type=int)
    semester = request.args.get('semester', default=1, type=int)
    cc_number_str = request.args.get('cc_number', default='all')

    selected_class = None
    if class_id:
        selected_class = Class.query.get(class_id)
    if not selected_class and classes:
        selected_class = classes[0]

    target_classes = [selected_class] if selected_class else classes

    missing_items = []
    affected_students_set = set()

    for c in target_classes:
        students = Student.query.filter_by(class_id=c.id).order_by(*Student.default_order()).all()
        exam_map = {
            eg.student_id: eg 
            for eg in SemesterExamGrade.query.filter_by(
                class_id=c.id,
                semester=semester,
                school_year_id=current_year.id
            ).all()
        }

        for st in students:
            ex = exam_map.get(st.id)
            ccs_to_check = [1, 2, 3, 4] if cc_number_str == 'all' else [int(cc_number_str)]
            for i in ccs_to_check:
                val = getattr(ex, f'cc{i}', None) if ex else None
                if val is None:
                    missing_items.append({
                        'student': st,
                        'classe': c,
                        'cc_number': i,
                        'exam_label': f'CC{i} - Contrôle Continu {i}'
                    })
                    affected_students_set.add(st.id)

    return render_template(
        'evaluation/makeup_manager.html',
        classes=classes,
        selected_class=selected_class,
        semester=semester,
        cc_number_str=cc_number_str,
        missing_items=missing_items,
        total_missing=len(missing_items),
        total_affected_students=len(affected_students_set),
        current_year=current_year
    )


@evaluation_bp.route('/make-up/print-sheet')
def makeup_print_sheet():
    """Feuille d'émargement officielle imprimable pour la session de rattrapage."""
    current_year = get_current_school_year()
    classes = Class.query.filter_by(school_year_id=current_year.id).order_by(Class.name).all() if current_year else Class.query.order_by(Class.name).all()

    class_id = request.args.get('class_id', type=int)
    semester = request.args.get('semester', default=1, type=int)
    cc_number_str = request.args.get('cc_number', default='all')
    room = request.args.get('room', default='Laboratoire Informatique')
    session_date = request.args.get('date', default=datetime.now().strftime('%d/%m/%Y'))

    selected_class = None
    if class_id:
        selected_class = Class.query.get(class_id)
    if not selected_class and classes:
        selected_class = classes[0]

    target_classes = [selected_class] if selected_class else classes

    missing_items = []
    for c in target_classes:
        students = Student.query.filter_by(class_id=c.id).order_by(*Student.default_order()).all()
        exam_map = {
            eg.student_id: eg 
            for eg in SemesterExamGrade.query.filter_by(
                class_id=c.id,
                semester=semester,
                school_year_id=current_year.id
            ).all()
        }

        for st in students:
            ex = exam_map.get(st.id)
            ccs_to_check = [1, 2, 3, 4] if cc_number_str == 'all' else [int(cc_number_str)]
            for i in ccs_to_check:
                val = getattr(ex, f'cc{i}', None) if ex else None
                if val is None:
                    missing_items.append({
                        'student': st,
                        'classe': c,
                        'cc_number': i,
                        'exam_label': f'CC{i}'
                    })

    return render_template(
        'evaluation/makeup_sheet.html',
        selected_class=selected_class,
        semester=semester,
        cc_number_str=cc_number_str,
        room=room,
        session_date=session_date,
        missing_items=missing_items,
        current_year=current_year
    )


@evaluation_bp.route('/api/save-makeup-grade', methods=['POST'])
def api_save_makeup_grade():
    """Sauvegarde rapide en AJAX de la note obtenue lors du rattrapage."""
    data = request.get_json() or {}
    student_id = data.get('student_id')
    class_id = data.get('class_id')
    semester = int(data.get('semester', 1))
    cc_number = int(data.get('cc_number', 1))
    grade_str = data.get('grade')

    if not student_id:
        return jsonify({'success': False, 'error': 'ID étudiant manquant'}), 400

    current_year = get_current_school_year()
    if not current_year:
        return jsonify({'success': False, 'error': 'Aucune année active'}), 400

    val = None
    if grade_str is not None and str(grade_str).strip() != '':
        try:
            val = round(max(0.0, min(20.0, float(str(grade_str).replace(',', '.')))), 2)
        except (ValueError, TypeError):
            val = None

    rec = SemesterExamGrade.query.filter_by(
        student_id=student_id,
        semester=semester,
        school_year_id=current_year.id
    ).first()

    if not rec:
        st = Student.query.get(student_id)
        rec = SemesterExamGrade(
            student_id=student_id,
            class_id=class_id or (st.class_id if st else 1),
            school_year_id=current_year.id,
            semester=semester
        )
        db.session.add(rec)

    setattr(rec, f'cc{cc_number}', val)
    db.session.commit()

    return jsonify({
        'success': True,
        'grade': val,
        'cc_average': rec.cc_average,
        'message': f'Note du CC{cc_number} enregistrée : {val}/20'
    })


@evaluation_bp.route('/settings')
def settings():
    """Interface de configuration des coefficients et formules de moyenne."""
    current_year = get_current_school_year()
    classes = Class.query.filter_by(school_year_id=current_year.id).order_by(Class.name).all() if current_year else Class.query.order_by(Class.name).all()

    class_id = request.args.get('class_id', default=0, type=int)
    selected_class = None
    if class_id > 0:
        selected_class = Class.query.get(class_id)

    # Vérifie si la classe a une configuration spécifique
    is_class_specific = False
    if selected_class:
        raw_class = AppSetting.get_value(f'eval_weights_class_{selected_class.id}', '')
        if raw_class:
            is_class_specific = True

    current_weights = get_evaluation_weights(selected_class.id if selected_class else None)
    global_weights = get_evaluation_weights(None)

    return render_template(
        'evaluation/settings.html',
        classes=classes,
        selected_class=selected_class,
        class_id=class_id,
        is_class_specific=is_class_specific,
        weights=current_weights,
        global_weights=global_weights
    )


@evaluation_bp.route('/settings/save', methods=['POST'])
def save_settings():
    """Enregistre les coefficients d'évaluation (globaux ou par classe)."""
    class_id = request.form.get('class_id', default=0, type=int)

    def parse_weight(key, default_val=1.0):
        val_str = request.form.get(key, '').strip().replace(',', '.')
        try:
            val = float(val_str)
            return max(0.0, min(10.0, val))
        except (ValueError, TypeError):
            return default_val

    weights_data = {
        'cc1': parse_weight('cc1', 1.0),
        'cc2': parse_weight('cc2', 1.0),
        'cc3': parse_weight('cc3', 1.0),
        'cc4': parse_weight('cc4', 1.0),
        'activity': parse_weight('activity', 1.0),
        'rounding': request.form.get('rounding', '0.01')
    }

    if class_id > 0:
        c = Class.query.get(class_id)
        c_name = c.name if c else f'ID {class_id}'
        AppSetting.set_value(f'eval_weights_class_{class_id}', json.dumps(weights_data), f"Coefficients d'évaluation pour classe {c_name}")
        flash(f"Coefficients personnalisés enregistrés avec succès pour la classe {c_name} !", 'success')
    else:
        AppSetting.set_value('eval_weights_global', json.dumps(weights_data), "Coefficients généraux d'évaluation (toutes les classes)")
        flash("Coefficients généraux enregistrés avec succès pour toutes les classes !", 'success')

    return redirect(url_for('evaluation.settings', class_id=class_id))


@evaluation_bp.route('/settings/reset', methods=['POST'])
def reset_settings():
    """Réinitialise les coefficients aux valeurs standard ou supprime la personnalisation par classe."""
    class_id = request.form.get('class_id', default=0, type=int)

    if class_id > 0:
        c = Class.query.get(class_id)
        c_name = c.name if c else f'ID {class_id}'
        setting = AppSetting.query.filter_by(key=f'eval_weights_class_{class_id}').first()
        if setting:
            db.session.delete(setting)
            db.session.commit()
        AppSetting.clear_cache()
        flash(f"La personnalisation pour la classe {c_name} a été supprimée. Elle utilise maintenant les coefficients généraux.", 'info')
    else:
        setting = AppSetting.query.filter_by(key='eval_weights_global').first()
        if setting:
            db.session.delete(setting)
            db.session.commit()
        AppSetting.clear_cache()
        flash("Les coefficients généraux ont été réinitialisés aux valeurs standards (1.0 par défaut).", 'info')

    return redirect(url_for('evaluation.settings', class_id=class_id))





